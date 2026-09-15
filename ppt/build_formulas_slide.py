# -*- coding: utf-8 -*-
"""FMECA-IPS 산출 공식 요약 -- 단일 슬라이드. ppt/build_deck.py와 동일한 팔레트/폰트를
그대로 써서 그 15슬라이드 덱에 바로 끼워넣을 수 있게 만든다."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DOMINANT = RGBColor(0x16, 0x30, 0x2B)
SECONDARY = RGBColor(0x2C, 0x6B, 0x73)
PANEL = RGBColor(0xF0, 0xF4, 0xF2)
LIGHT_LINE = RGBColor(0xDD, 0xE3, 0xDF)
INK = RGBColor(0x1A, 0x24, 0x22)
INK_DIM = RGBColor(0x55, 0x63, 0x5E)
INK_FAINT = RGBColor(0x8B, 0x96, 0x8F)
AMBER = RGBColor(0xB5, 0x65, 0x1D)
AMBER_BG = RGBColor(0xF7, 0xEB, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def set_run(run, text, size, color, bold=False, font=BODY_FONT, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             wrap=True, margin=0.0, line_spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    if isinstance(lines[0], tuple):
        lines = [[ln] for ln in lines]
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for run_spec in para:
            text, size, color = run_spec[0], run_spec[1], run_spec[2]
            bold = run_spec[3] if len(run_spec) > 3 else False
            font = run_spec[4] if len(run_spec) > 4 else BODY_FONT
            italic = run_spec[5] if len(run_spec) > 5 else False
            r = p.add_run()
            set_run(r, text, size, color, bold, font, italic)
    return tb


def add_rect(slide, x, y, w, h, fill, line_color=None, line_w=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = line_w or Pt(0.75)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def formula_card(slide, x, y, w, h, title, title_color, lines):
    """lines: list of (formula_text, note_text_or_None)."""
    add_rect(slide, x, y, w, h, PANEL, line_color=LIGHT_LINE, line_w=Pt(0.75))
    add_text(slide, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.3),
              [[(title, 14, title_color, True, HEAD_FONT)]])
    yy = y + Inches(0.52)
    for formula, note in lines:
        paras = [[(formula, 11.3, INK, True, MONO_FONT)]]
        if note:
            paras.append([(note, 9, INK_DIM, False, BODY_FONT, True)])
        add_text(slide, x + Inches(0.18), yy, w - Inches(0.36), Inches(0.62), paras, line_spacing=1.05)
        yy += Inches(0.44) if note else Inches(0.3)


def rule_card(slide, x, y, w, h, title, title_color, rule_lines):
    add_rect(slide, x, y, w, h, WHITE, line_color=LIGHT_LINE, line_w=Pt(0.75))
    add_rect(slide, x, y, w, Inches(0.34), title_color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    add_rect(slide, x, y + Inches(0.17), w, Inches(0.17), title_color, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, x, y + Inches(0.03), w, Inches(0.3), [[(title, 12.5, WHITE, True, HEAD_FONT)]], align=PP_ALIGN.CENTER)
    yy = y + Inches(0.46)
    for text in rule_lines:
        add_text(slide, x + Inches(0.14), yy, w - Inches(0.28), Inches(0.5),
                  [[(text, 9.7, INK, False, BODY_FONT)]], line_spacing=1.08)
        yy += Inches(0.42)


# ===========================================================================
s = new_slide()
add_text(s, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.32),
          [[("산출 공식 요약", 12, SECONDARY, True, BODY_FONT)]])
add_text(s, Inches(0.53), Inches(0.55), Inches(12.2), Inches(0.55),
          [[("FMECA O·S·D·위험지수 산출식과 IPS 연계 판정 규칙", 24, INK, True, HEAD_FONT)]])

# ---- row 1: O / S / D / RPN,RI ----
row1_y = Inches(1.35)
row1_h = Inches(2.55)
card_w = Inches(3.0)
gap = Inches(0.2)
x0 = Inches(0.55)

formula_card(s, x0, row1_y, card_w, row1_h, "발생도 (O)", SECONDARY, [
    ("rate = λ × 1000", "λ=수리건수÷노출시간"),
    ("O등급 = bin(log₁₀(rate))", "관측범위 10구간 균등분할"),
    ("CI = χ²(0.025,2k)/2,", "χ²(0.975,2(k+1))/2"),
])

formula_card(s, x0 + (card_w+gap)*1, row1_y, card_w, row1_h, "심각도 (S, 근사)", AMBER, [
    ("guide word = MORE/LESS", "부호부 z 평균의 방향"),
    ("오차율% = |rate−기준선| /", "기준선 × 100"),
    ("S = round(1+9×S_norm)", "결과심각도 기록 없어 대리지표"),
])

formula_card(s, x0 + (card_w+gap)*2, row1_y, card_w, row1_h, "검출도 (D)", SECONDARY, [
    ("탐지력 = 0.6×탐지율 +", "0.4×min(리드타임/60,1)"),
    ("D = clip(round(10−9×탐지력)", "1,10)"),
    ("p = Binomial(n,p_null)", "이항검정, p<0.05면 유의"),
])

formula_card(s, x0 + (card_w+gap)*3, row1_y, card_w, row1_h, "위험지수", AMBER, [
    ("RPN = O × S × D", "곱셈식, 극단값에 민감"),
    ("RI = 0.4O+0.4S+0.2D", "가중합산식, 완만하게 반영"),
    ("두 지수 순위 비교", "왜곡 여부 상호 검증"),
])

# ---- row 2: IPS linkage rules ----
row2_y = Inches(4.15)
add_text(s, x0, row2_y, Inches(12.2), Inches(0.3), [[("IPS 연계 판정 규칙", 14, INK, True, HEAD_FONT)]])
row2_y2 = row2_y + Inches(0.42)
row2_h = Inches(2.55)
rc_w = Inches(2.36)
rgap = Inches(0.1)

rule_card(s, x0, row2_y2, rc_w, row2_h, "RCM", SECONDARY, [
    "유의(p<.05)&탐지율≥50%",
    "→ CBM",
    "else S≥5 → Hard-Time",
    "else → RTF",
])
rule_card(s, x0+(rc_w+rgap)*1, row2_y2, rc_w, row2_h, "LORA", SECONDARY, [
    "O≥8",
    "→ 창정비 우선상향+",
    "예비품 전진배치",
    "else → 표준 2-Level",
])
rule_card(s, x0+(rc_w+rgap)*2, row2_y2, rc_w, row2_h, "보급소요", SECONDARY, [
    "annual_demand=",
    "rate×fleet×가동일수",
    "reorder_point=",
    "리드타임수요+안전재고",
])
rule_card(s, x0+(rc_w+rgap)*3, row2_y2, rc_w, row2_h, "IETM", SECONDARY, [
    "유의&리드타임 존재",
    "→ SLA=리드타임×0.5",
    "else → 정기점검",
    "강화(SLA 없음)",
])
rule_card(s, x0+(rc_w+rgap)*4, row2_y2, rc_w, row2_h, "ECP", AMBER, [
    "RPN 1위 / O=10 /",
    "비유의 중 하나",
    "→ 설계개선 후보",
    "else → 현행유지",
])

add_text(s, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.3),
          [[("전체 수식의 \"왜 이렇게 계산하는가\" 논리 설명은 doc/FMECA_IPS_공식정리.docx 참고 · 실측 원천: SCANIA Component X (CC BY 4.0)", 9, INK_FAINT, False, BODY_FONT)]])

prs.save("FMECA_IPS_공식요약_1p.pptx")
print("saved")
