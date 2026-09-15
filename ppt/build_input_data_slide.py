# -*- coding: utf-8 -*-
"""입력 데이터 구조 -- 단일 슬라이드. build_formulas_slide.py의 짝(산출식 앞에 오는 원시
데이터 구조)이며 같은 팔레트/폰트를 쓴다. datasets/scania_component_x/config.yaml이
실제로 채우는 형태를 그대로 시각화한다 -- DATASET_FORMAT.md의 3-CSV + config.yaml 계약."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

DOMINANT = RGBColor(0x16, 0x30, 0x2B)
SECONDARY = RGBColor(0x2C, 0x6B, 0x73)
PANEL = RGBColor(0xF0, 0xF4, 0xF2)
LIGHT_LINE = RGBColor(0xDD, 0xE3, 0xDF)
INK = RGBColor(0x1A, 0x24, 0x22)
INK_DIM = RGBColor(0x55, 0x63, 0x5E)
INK_FAINT = RGBColor(0x8B, 0x96, 0x8F)
AMBER = RGBColor(0xB5, 0x65, 0x1D)
AMBER_BG = RGBColor(0xF7, 0xEB, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def set_run(run, text, size, color, bold=False, font=BODY_FONT, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             wrap=True, margin=0.0, line_spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    if isinstance(lines[0], tuple):
        lines = [[ln] for ln in lines]
    for i, para in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for run_spec in para:
            text, size, color = run_spec[0], run_spec[1], run_spec[2]
            bold = run_spec[3] if len(run_spec) > 3 else False
            font = run_spec[4] if len(run_spec) > 4 else BODY_FONT
            italic = run_spec[5] if len(run_spec) > 5 else False
            r = p.add_run()
            set_run(r, text, size, color, bold, font, italic)
    return tb


def add_rect(slide, x, y, w, h, fill, line_color=None, line_w=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = line_w or Pt(0.75)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def add_arrow(slide, x1, y1, x2, y2, color=INK_DIM, w=Pt(1.25), dashed=False):
    """Straight connector line, drawn as a thin filled rectangle instead of a
    genuine Connector autoshape -- python-pptx's add_connector produces XML
    that PowerPoint's own (strict) reader rejects as corrupt even though
    python-pptx and generic OOXML validators accept it fine."""
    thickness = Emu(int(w))
    if x1 == x2:
        rx, ry, rw, rh = x1 - thickness/2, min(y1, y2), thickness, abs(y2 - y1)
    else:
        rx, ry, rw, rh = min(x1, x2), y1 - thickness/2, abs(x2 - x1), thickness
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, ry, max(rw, Emu(1)), max(rh, Emu(1)))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def file_card(slide, x, y, w, h, title, subtitle, fields, accent):
    add_rect(slide, x, y, w, h, WHITE, line_color=LIGHT_LINE, line_w=Pt(0.75))
    add_rect(slide, x, y, w, Inches(0.42), accent, radius=0.12)
    add_rect(slide, x, y + Inches(0.21), w, Inches(0.21), accent, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, x + Inches(0.12), y + Inches(0.03), w - Inches(0.24), Inches(0.24),
              [[(title, 13, WHITE, True, MONO_FONT)]])
    add_text(slide, x + Inches(0.12), y + Inches(0.24), w - Inches(0.24), Inches(0.18),
              [[(subtitle, 8.5, WHITE, False, BODY_FONT)]])
    yy = y + Inches(0.56)
    for name, desc in fields:
        add_text(slide, x + Inches(0.14), yy, w - Inches(0.28), Inches(0.28),
                  [[(name + "  ", 10, INK, True, MONO_FONT), (desc, 8.7, INK_DIM, False, BODY_FONT, True)]],
                  line_spacing=1.0)
        yy += Inches(0.315)


# ===========================================================================
s = new_slide()
add_text(s, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.32),
          [[("입력 데이터 구조", 12, SECONDARY, True, BODY_FONT)]])
add_text(s, Inches(0.53), Inches(0.55), Inches(12.2), Inches(0.55),
          [[("자산ID로 연결되는 CSV 3종 + config.yaml — 데이터셋이 바뀌어도 형식은 고정", 22, INK, True, HEAD_FONT)]])

# ---- 3 file cards ----
fy = Inches(1.35)
fh = Inches(2.55)
fw = Inches(3.75)
fx0 = Inches(0.55)
gap = Inches(0.25)

file_card(s, fx0, fy, fw, fh, "events.csv", "자산별 1행 — 노출시간·고장라벨", [
    ("asset_id", "자산 식별자"),
    ("exposure_time", "노출시간(숫자)"),
    ("failure_label", "고장라벨(0/1)"),
], SECONDARY)

file_card(s, fx0 + fw + gap, fy, fw, fh, "readouts.csv", "자산×시점 1행 — 시계열 계측값", [
    ("asset_id", "자산 식별자"),
    ("readout_time", "판독 시점"),
    ("counters[]", "누적형 계측 카운터 1개+"),
], SECONDARY)

file_card(s, fx0 + (fw + gap) * 2, fy, fw, fh, "groups.csv", "자산별 1행 — FMECA 행 그룹핑", [
    ("asset_id", "자산 식별자"),
    ("group_col", "FMECA 행 구분 그룹"),
    ("bonus_group_col", "선택: 보조 드릴다운"),
], SECONDARY)

# join arrows -> config.yaml band
join_y = fy + fh + Inches(0.05)
for i in range(3):
    cx = fx0 + (fw + gap) * i + fw/2
    add_arrow(s, cx, join_y, cx, join_y + Inches(0.32))
add_text(s, fx0, join_y + Inches(0.06), fw*3 + gap*2, Inches(0.24),
          [[("asset_id 로 조인", 10, INK_FAINT, True, BODY_FONT, True)]], align=PP_ALIGN.CENTER)

# ---- config.yaml band ----
cfg_y = fy + fh + Inches(0.55)
cfg_h = Inches(1.55)
add_rect(s, fx0, cfg_y, fw*3 + gap*2, cfg_h, AMBER_BG, line_color=AMBER, line_w=Pt(1))
add_text(s, fx0 + Inches(0.2), cfg_y + Inches(0.12), Inches(4), Inches(0.3),
          [[("config.yaml", 14, AMBER, True, MONO_FONT)]])
cfg_cols = [
    ["columns:", "3개 CSV의 실제 컬럼명 →", "의미역할 매핑"],
    ["groups: / params:", "그룹라벨 · 임계값", "(z_thresh, lookback 등)"],
    ["assumptions:", "IPS 가정치 — value +", "rationale(근거) 필수"],
    ["content:", "웹앱 서술문(매핑표·", "공백표·결과해석)"],
]
cw = (fw*3 + gap*2 - Inches(0.4)) / 4
for i, lines in enumerate(cfg_cols):
    xx = fx0 + Inches(0.2) + cw * i
    add_text(s, xx, cfg_y + Inches(0.55), cw - Inches(0.1), Inches(0.9),
              [[(lines[0], 10.5, INK, True, MONO_FONT)]] + [[(ln, 8.8, INK_DIM, False, BODY_FONT)] for ln in lines[1:]],
              line_spacing=1.05)

# ---- arrow down to pipeline ----
pipe_y = cfg_y + cfg_h + Inches(0.05)
add_arrow(s, fx0 + (fw*3+gap*2)/2, pipe_y, fx0 + (fw*3+gap*2)/2, pipe_y + Inches(0.28))

add_rect(s, fx0, pipe_y + Inches(0.3), fw*3 + gap*2, Inches(0.85), PANEL, line_color=LIGHT_LINE, line_w=Pt(0.75))
add_text(s, fx0 + Inches(0.2), pipe_y + Inches(0.42), fw*3+gap*2 - Inches(0.4), Inches(0.6),
          [[("6단계 정제 (동기화→드리프트제거→구간분할→결측/이상치→코드정합→라벨정의)", 11.5, INK, True, BODY_FONT)],
           [("→ 자산별 rate(t)·z(t)·anomaly(t) 계산 → O, S, D 산출로 이어짐 (별도 슬라이드: 산출 공식 요약)", 10, INK_DIM, False, BODY_FONT, True)]],
          align=PP_ALIGN.CENTER, line_spacing=1.15)

add_text(s, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.3),
          [[("스키마 상세는 DATASET_FORMAT.md 참고 · 현재 예시: SCANIA Component X (23,550대, CC BY 4.0) — events=train_tte.csv, readouts=train_operational_readouts.csv, groups=train_specifications.csv", 8.7, INK_FAINT, False, BODY_FONT)]])

prs.save("FMECA_IPS_입력데이터구조_1p.pptx")
print("saved")
